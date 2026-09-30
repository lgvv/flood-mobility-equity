#!/usr/bin/env python3
"""Extract location and facility class from the public HIRA quarterly workbook ZIP."""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import io
import json
from pathlib import Path
import zipfile

import pandas as pd
from openpyxl import load_workbook


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def rows_from_excel(payload: bytes, source: str):
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        records = sheet.iter_rows(values_only=True)
        header = [str(x).strip() if x is not None else "" for x in next(records)]
        required = ["암호화요양기호", "종별코드명", "시도코드명", "시군구코드명", "좌표(X)", "좌표(Y)"]
        missing = set(required) - set(header)
        if missing:
            raise ValueError(f"{source}: missing columns {missing}")
        indices = [header.index(name) for name in required]
        for row in records:
            yield (source, *(row[i] if i < len(row) else None for i in indices))
    finally:
        workbook.close()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", type=Path, default=Path("data/raw/hira_facilities_202606.zip"))
    ap.add_argument("--out", type=Path, default=Path("data/processed"))
    args = ap.parse_args()
    candidates = []
    with zipfile.ZipFile(args.source) as archive:
        for info in archive.infolist():
            name = info.filename.split("/")[-1]
            if name.startswith(("1.", "2.")) and name.endswith(".xlsx"):
                candidates.extend(rows_from_excel(archive.read(info), name))
    columns = ["source_workbook", "facility_id", "facility_type", "province", "district", "longitude", "latitude"]
    raw = pd.DataFrame(candidates, columns=columns)
    raw["longitude"] = pd.to_numeric(raw["longitude"], errors="coerce")
    raw["latitude"] = pd.to_numeric(raw["latitude"], errors="coerce")
    valid = raw.longitude.between(124, 132) & raw.latitude.between(33, 39)
    blank_id = raw.facility_id.isna() | raw.facility_id.astype(str).str.strip().eq("")
    clean = raw.loc[valid & ~blank_id].drop_duplicates("facility_id").copy()
    clean = clean.sort_values("facility_id")
    args.out.mkdir(parents=True, exist_ok=True)
    clean.to_csv(args.out / "hira_medical_sites_202606.csv", index=False, encoding="utf-8-sig")
    receipt = {
        "source_url": "https://opendata.hira.or.kr/op/opc/selectOpenData.do?sno=11925",
        "source_reference_month": "2026-06",
        "source_sha256": sha256(args.source),
        "raw_rows": len(raw),
        "valid_coordinate_rows": int(valid.sum()),
        "invalid_coordinate_rows": int((~valid).sum()),
        "blank_id_rows": int(blank_id.sum()),
        "duplicate_id_rows_after_validity": int((valid & ~blank_id).sum() - len(clean)),
        "retained_rows": len(clean),
        "facility_types": dict(Counter(clean.facility_type.fillna("MISSING"))),
        "coordinates": "Source X/Y values were inspected and interpreted as longitude/latitude WGS84 degrees.",
        "medical_class_warning": "Facility type is not an emergency designation; do not call all hospitals emergency facilities.",
    }
    (args.out / "hira_medical_sites_202606_qa.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(receipt, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
